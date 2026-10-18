"""Document parsing tests using generated, non-private documents."""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from bot.services.documents import MAX_DOCUMENT_BYTES, extract_document_text


def make_pdf(*pages: str) -> bytes:
    """Create small text PDFs without adding a PDF-generation dependency."""
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(
            f"BT /F1 12 Tf 10 10 Td ({text}) Tj ET".encode("ascii")
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_preserves_text_from_all_pages() -> None:
    text = extract_document_text(
        make_pdf("First advice", "Second advice"), "a.PDF"
    )

    assert "First advice" in text
    assert "Second advice" in text
    assert text.index("First advice") < text.index("Second advice")


def test_docx_preserves_paragraph_and_table_order() -> None:
    document = Document()
    document.add_paragraph("Початок")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Місце"
    table.cell(0, 1).text = "Порада"
    document.add_paragraph("Завершення")
    output = BytesIO()
    document.save(output)

    assert extract_document_text(output.getvalue(), "a.docx") == (
        "Початок\nМісце\tПорада\nЗавершення"
    )


def test_docx_preserves_nested_table_text_and_cell_block_order() -> None:
    document = Document()
    document.add_paragraph("Document before")
    outer = document.add_table(rows=1, cols=2)
    cell = outer.cell(0, 0)
    cell.text = "Outer before"
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested left"
    nested.cell(0, 1).text = "Nested right"
    cell.paragraphs[-1].text = "Outer after"
    outer.cell(0, 1).text = "Second outer cell"
    document.add_paragraph("Document after")
    output = BytesIO()
    document.save(output)

    assert extract_document_text(output.getvalue(), "nested.docx") == (
        "Document before\nOuter before\nNested left\tNested right\n"
        "Outer after\tSecond outer cell\nDocument after"
    )


@pytest.mark.parametrize("filename", ["a.pdf", "a.docx", "a.txt"])
def test_empty_file_is_rejected(filename: str) -> None:
    with pytest.raises(ValueError, match="empty"):
        extract_document_text(b"", filename)


@pytest.mark.parametrize("filename", ["a.pdf", "a.docx"])
def test_damaged_file_is_rejected(filename: str) -> None:
    with pytest.raises(ValueError, match="damaged"):
        extract_document_text(b"not a document", filename)


def test_pdf_without_text_is_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)

    with pytest.raises(ValueError, match="no extractable text"):
        extract_document_text(output.getvalue(), "scan.pdf")


def test_empty_docx_is_rejected() -> None:
    output = BytesIO()
    Document().save(output)

    with pytest.raises(ValueError, match="no extractable text"):
        extract_document_text(output.getvalue(), "empty.docx")


def test_encrypted_pdf_is_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("private-password")
    output = BytesIO()
    writer.write(output)

    with pytest.raises(ValueError, match="encrypted"):
        extract_document_text(output.getvalue(), "encrypted.pdf")


def test_unsupported_format_is_rejected() -> None:
    with pytest.raises(ValueError, match="Supported"):
        extract_document_text(b"legacy document", "old.doc")


def test_oversized_document_is_rejected() -> None:
    with pytest.raises(ValueError, match="size limit"):
        extract_document_text(b"x" * (MAX_DOCUMENT_BYTES + 1), "large.txt")


@pytest.mark.parametrize(
    ("content", "filename", "expected"),
    [
        pytest.param(
            "Правила лісу".encode("utf-8"),
            "rules.TXT",
            "Правила лісу",
            id="ukrainian-uppercase-extension",
        ),
        pytest.param(
            b"\xef\xbb\xbf" + " \tПравила лісу\r\n".encode("utf-8"),
            "rules.txt",
            "Правила лісу",
            id="bom-and-outer-whitespace",
        ),
        pytest.param(
            "  Перше\r\n\r\nДруге\tпункт \t".encode("utf-8"),
            "rules.txt",
            "Перше\r\n\r\nДруге\tпункт",
            id="internal-formatting",
        ),
    ],
)
def test_txt_accepts_utf8_and_preserves_internal_formatting(
    content: bytes, filename: str, expected: str
) -> None:
    assert extract_document_text(content, filename) == expected


@pytest.mark.parametrize(
    "content",
    [
        pytest.param(b" \t\r\n", id="whitespace-only"),
        pytest.param(b"\xef\xbb\xbf", id="bom-only"),
        pytest.param(b"\xef\xbb\xbf \t\r\n", id="bom-and-whitespace"),
    ],
)
def test_txt_without_text_is_rejected(content: bytes) -> None:
    with pytest.raises(ValueError, match="empty"):
        extract_document_text(content, "rules.txt")


def test_txt_invalid_utf8_after_valid_prefix_is_rejected() -> None:
    with pytest.raises(ValueError, match="UTF-8"):
        extract_document_text(b"Valid prefix\n\xff", "rules.txt")
